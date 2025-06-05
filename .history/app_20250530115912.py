from flask import Flask, request, jsonify, send_from_directory
import os
import uuid
import json
from datetime import datetime
import pandas as pd
import docx
import pytesseract
from pdf2image import convert_from_path
from PyPDF2 import PdfReader
from werkzeug.utils import secure_filename
from fpdf import FPDF
import spacy

nlp = spacy.load("en_core_web_sm")

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
ANALYSIS_FOLDER = 'analysis'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(ANALYSIS_FOLDER, exist_ok=True)

ALLOWED_EXTENSIONS = {'pdf', 'docx', 'png', 'jpg', 'jpeg'}
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['ANALYSIS_FOLDER'] = ANALYSIS_FOLDER

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_docx(filepath):
    doc = docx.Document(filepath)
    content = []
    for para in doc.paragraphs:
        style = para.style.name
        if para.text.strip():
            content.append((style, para.text.strip()))
    return content

def extract_text_pdf(filepath):
    try:
        reader = PdfReader(filepath)
        text = ''
        for page in reader.pages:
            text += page.extract_text() or ''
        if not text.strip():
            # OCR fallback
            images = convert_from_path(filepath)
            for image in images:
                text += pytesseract.image_to_string(image)
        return [("BodyText", text.strip())]
    except Exception as e:
        return [("Error", str(e))]

def extract_text_image(filepath):
    try:
        text = pytesseract.image_to_string(filepath)
        return [("ImageText", text.strip())]
    except Exception as e:
        return [("Error", str(e))]

def summarize_text(text):
    doc = nlp(text)
    sentences = list(doc.sents)
    summary = ' '.join([str(s) for s in sentences[:3]])
    return summary

def generate_pdf_report(file_id, extracted):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Document Summary Report", ln=True, align='C')
    pdf.ln(10)
    for style, content in extracted:
        pdf.set_font("Arial", size=10 if style == 'BodyText' else 12, style='B' if "Heading" in style else '')
        lines = content.split('\n')
        for line in lines:
            pdf.multi_cell(0, 10, txt=line.strip())
    pdf_path = os.path.join(app.config['ANALYSIS_FOLDER'], f"{file_id}_summary.pdf")
    pdf.output(pdf_path)
    return pdf_path

@app.route('/api/analyze', methods=['POST'])
def analyze_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400

    if not allowed_file(file.filename):
        return jsonify({"error": "Unsupported file type"}), 400

    filename = secure_filename(file.filename)
    file_id = str(uuid.uuid4())
    saved_path = os.path.join(app.config['UPLOAD_FOLDER'], f"{file_id}_{filename}")
    file.save(saved_path)

    ext = filename.rsplit('.', 1)[1].lower()
    if ext == 'docx':
        extracted = extract_text_docx(saved_path)
    elif ext == 'pdf':
        extracted = extract_text_pdf(saved_path)
    elif ext in ['png', 'jpg', 'jpeg']:
        extracted = extract_text_image(saved_path)
    else:
        return jsonify({"error": "Unsupported format"}), 400

    full_text = '\n'.join([text for _, text in extracted])
    summary = summarize_text(full_text)
    pdf_path = generate_pdf_report(file_id, extracted)

    analysis_log = {
        "file_id": file_id,
        "filename": filename,
        "summary": summary,
        "extracted_sections": extracted,
        "report_pdf": f"/api/download/report/{file_id}_summary.pdf",
        "timestamp": datetime.utcnow().isoformat()
    }

    log_path = os.path.join(app.config['ANALYSIS_FOLDER'], f"{file_id}_log.json")
    with open(log_path, 'w') as f:
        json.dump(analysis_log, f, indent=4)

    return jsonify(analysis_log)

@app.route('/api/download/report/<filename>', methods=['GET'])
def download_report(filename):
    return send_from_directory(app.config['ANALYSIS_FOLDER'], filename, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
