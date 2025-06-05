from flask import Flask, request, jsonify, send_from_directory
import pandas as pd
import os
import uuid
import joblib
from werkzeug.utils import secure_filename
import json
from datetime import datetime

# File parsing
import docx
from PyPDF2 import PdfReader
from pdf2image import convert_from_path
import pytesseract
from werkzeug.utils import secure_filename

# ML
from sklearn.model_selection import train_test_split
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import accuracy_score

# NLP
import spacy
from transformers import pipeline

# PDF generation
from fpdf import FPDF

# --- Setup ---
app = Flask(__name__)

UPLOAD_FOLDER = 'uploads'
ANALYSIS_FOLDER = 'analysis'
MODEL_FOLDER = 'models'
PREVIEW_FOLDER = 'previews'
LOGS_FOLDER = 'logs'

for folder in [UPLOAD_FOLDER, MODEL_FOLDER, PREVIEW_FOLDER, LOGS_FOLDER]:
    os.makedirs(folder, exist_ok=True)

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MODEL_FOLDER'] = MODEL_FOLDER
app.config['PREVIEW_FOLDER'] = PREVIEW_FOLDER
app.config['ANALYSIS_FOLDER'] = ANALYSIS_FOLDER


ALLOWED_EXTENSIONS = {'csv', 'txt', 'pdf', 'docx', 'png', 'jpg', 'jpeg'}

# Load NLP models
nlp = spacy.load("en_core_web_sm")
summarizer = pipeline("summarization", model="sshleifer/distilbart-cnn-12-6")

# --- Utility Functions ---
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def parse_file(filepath):
    ext = filepath.rsplit('.', 1)[1].lower()
    if ext == 'csv':
        return pd.read_csv(filepath)
    elif ext == 'txt':
        with open(filepath, 'r') as f:
            lines = f.readlines()
        return pd.DataFrame([line.strip().split(',') for line in lines[1:]], columns=lines[0].strip().split(','))
    elif ext == 'docx':
        raise NotImplementedError("Structured DOCX parsing not implemented.")
    elif ext == 'pdf':
        raise NotImplementedError("Structured PDF parsing not implemented.")
    else:
        raise ValueError("Unsupported file type")

def extract_sections_by_style(doc):
    sections = {}
    current_heading = None
    current_content = []

    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            if current_heading:
                sections[current_heading] = ' '.join(current_content).strip()
            current_heading = para.text.strip()
            current_content = []
        elif current_heading:
            current_content.append(para.text.strip())

    if current_heading and current_content:
        sections[current_heading] = ' '.join(current_content).strip()

    return sections

def generate_pdf_summary(data, output_path):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Project Analysis Summary", ln=True, align="C")
    pdf.ln(10)

    for heading, content in data.items():
        pdf.set_font("Arial", "B", 14)
        pdf.multi_cell(0, 10, heading)
        pdf.set_font("Arial", "", 12)
        pdf.multi_cell(0, 10, content)
        pdf.ln(5)

    pdf.output(output_path)

def save_training_log(file_id, original_filename, model_type, accuracy, dataset_shape):
    log_data = {
        "file_id": file_id,
        "original_filename": original_filename,
        "model_type": model_type,
        "accuracy": accuracy,
        "num_rows": dataset_shape[0],
        "num_columns": dataset_shape[1],
        "timestamp": datetime.utcnow().isoformat() + 'Z'
    }

    log_path = os.path.join(LOGS_FOLDER, f"{file_id}_log.json")
    with open(log_path, 'w') as f:
        json.dump(log_data, f, indent=4)

# --- Routes ---
@app.route('/api/train', methods=['POST'])
def train_model():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400

    file = request.files['file']
    model_type = request.form.get('model_type', 'random_forest')

    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400

    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_id = str(uuid.uuid4())
        saved_filename = f"{file_id}_{filename}"
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], saved_filename)
        file.save(filepath)

        try:
            df = parse_file(filepath)
            if df.shape[0] < 5 or df.shape[1] < 2:
                return jsonify({"error": "Insufficient data for training"}), 400

            preview_path = os.path.join(PREVIEW_FOLDER, f"{file_id}_preview.csv")
            df.head(10).to_csv(preview_path, index=False)

            df = df.apply(pd.to_numeric, errors='ignore')
            X = df.iloc[:, :-1]
            y = df.iloc[:, -1]

            X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2)

            if model_type == "random_forest":
                model = RandomForestClassifier()
            elif model_type == "logistic_regression":
                model = LogisticRegression(max_iter=1000)
            else:
                return jsonify({"error": "Unsupported model type"}), 400

            model.fit(X_train, y_train)
            y_pred = model.predict(X_test)
            accuracy = accuracy_score(y_test, y_pred)

            model_path = os.path.join(MODEL_FOLDER, f"{file_id}_{model_type}.pkl")
            joblib.dump(model, model_path)

            save_training_log(file_id, filename, model_type, accuracy, df.shape)

            return jsonify({
                "message": "Model trained successfully",
                "accuracy": accuracy,
                "model_download_url": f"/api/download/model/{file_id}_{model_type}.pkl",
                "preview_download_url": f"/api/download/preview/{file_id}_preview.csv"
            })

        except NotImplementedError as e:
            return jsonify({"error": str(e)}), 501
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    return jsonify({"error": "Unsupported file type"}), 400

@app.route('/api/analyze-doc', methods=['POST'])
def analyze_doc():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400

    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_id = str(uuid.uuid4())
        saved_filename = f"{file_id}_{filename}"
        filepath = os.path.join(UPLOAD_FOLDER, saved_filename)
        file.save(filepath)

        try:
            ext = filename.rsplit('.', 1)[1].lower()

            if ext != 'docx':
                return jsonify({"error": "Only .docx supported for heading-style NLP analysis"}), 400

            doc = docx.Document(filepath)
            sections = extract_sections_by_style(doc)

            analysis = {}
            for heading, content in sections.items():
                if len(content) > 30:
                    summary = summarizer(content, max_length=100, min_length=30, do_sample=False)[0]['summary_text']
                else:
                    summary = content
                analysis[heading] = summary

            json_path = os.path.join(LOGS_FOLDER, f"{file_id}_analysis.json")
            with open(json_path, 'w') as f:
                json.dump(analysis, f, indent=4)

            pdf_path = os.path.join(LOGS_FOLDER, f"{file_id}_summary.pdf")
            generate_pdf_summary(analysis, pdf_path)

            return jsonify({
                "message": "Analysis completed",
                "sections": list(analysis.keys()),
                "json_log_url": f"/api/download/log/{file_id}_analysis.json",
                "pdf_summary_url": f"/api/download/log/{file_id}_summary.pdf"
            })

        except Exception as e:
            return jsonify({"error": str(e)}), 500

    return jsonify({"error": "Unsupported file type"}), 400

@app.route('/api/download/model/<filename>', methods=['GET'])
def download_model(filename):
    return send_from_directory(MODEL_FOLDER, filename, as_attachment=True)

@app.route('/api/download/preview/<filename>', methods=['GET'])
def download_preview(filename):
    return send_from_directory(PREVIEW_FOLDER, filename, as_attachment=True)

@app.route('/api/download/log/<filename>', methods=['GET'])
def download_log(filename):
    return send_from_directory(LOGS_FOLDER, filename, as_attachment=True)

# --- Main ---
if __name__ == '__main__':
    app.run(debug=True)
